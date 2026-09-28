import io
import unittest

from openpyxl import load_workbook

from influencer_core import InfluencerConfig, collect_influencers, evaluate_apify_profile
from influencer_exporter import build_influencer_excel
from version import APP_VERSION


def config(**changes):
    values = {
        "country": "Vietnam",
        "category": "뷰티·화장품",
        "keywords": ["beauty", "skincare", "makeup"],
        "target_count": 10,
        "max_queries": 4,
        "request_delay": 0,
        "use_apify_verification": True,
        "use_apify_discovery": True,
        "use_web_fallback": False,
    }
    values.update(changes)
    return InfluencerConfig(**values)


def profile(handle, full_name, biography, category="Blogger", followers=5000, captions=None, **extra):
    row = {
        "inputUrl": f"https://www.instagram.com/{handle}/",
        "url": f"https://www.instagram.com/{handle}/",
        "username": handle,
        "fullName": full_name,
        "biography": biography,
        "businessCategoryName": category,
        "followersCount": followers,
        "latestPosts": [{"caption": text} for text in (captions or [])],
    }
    row.update(extra)
    return row


class V2301Tests(unittest.TestCase):
    def test_version(self):
        self.assertEqual(APP_VERSION, "2.30.3")

    def test_explicit_vietnam_nha_trang_profile_is_high_confidence(self):
        row = profile(
            "morelly_angelina",
            "beauty influencer | UGC creator | lifestyle blogger Vietnam 🇻🇳",
            "Angelina Morelly 📍Vietnam, Nha Trang women blog about beauty through health & love",
            captions=["My skincare and beauty routine"] * 6,
            followers=1287,
        )
        result = evaluate_apify_profile(row, config())
        self.assertEqual(result["country"], "Vietnam")
        self.assertEqual(result["country_confidence"], "높음")
        self.assertTrue(result["eligible"])

    def test_direct_beauty_profile_wording_confirms_category(self):
        row = profile(
            "trinhmeow.vlogger",
            "Trinh Meow",
            "📸 Beauty Blogger 🧬Skinfluencer · Based in Saigon",
            category="Personal blog",
            captions=["Một ngày bình yên tại Sài Gòn"] * 12,
            followers=17562,
        )
        result = evaluate_apify_profile(row, config())
        self.assertEqual(result["category_relevance"], "확인")
        self.assertTrue(result["eligible"])
        self.assertIn("직접표현", result["category_evidence"])

    def test_observed_brand_accounts_are_excluded_after_profile_verification(self):
        for handle, name in (
            ("mynnetwork", "MYN Network"),
            ("rodinshovietnam", "Rodin Sho"),
            ("leve.co.ltd", "Leve Cosmetics"),
            ("magamakeupbrushes", "MAGA Makeup Brushes"),
            ("m.u.i_perfume", "M.U.I perfume"),
            ("newbeautynetworkvn", "New Beauty Network"),
            ("glow_cosmetics", "Glow Cosmetics"),
            ("sample.co.ltd", "Sample Beauty"),
        ):
            with self.subTest(handle=handle):
                row = profile(
                    handle,
                    name,
                    "Vietnam beauty and skincare products",
                    category="Product/service",
                    captions=["Sản phẩm chăm sóc da và mỹ phẩm"] * 12,
                    followers=5000,
                    isBusinessAccount=True,
                )
                result = evaluate_apify_profile(row, config())
                self.assertEqual(result["status"], "제외_기업계정")

    def test_medium_country_confidence_is_final_grade_b_not_a(self):
        discovery = {
            "username": "hufureview",
            "fullName": "Review Mỹ phẩm và Tips Làm đẹp",
            "biography": "Tips làm đẹp và review mỹ phẩm",
            "businessCategoryName": "Reel creator",
            "followersCount": 17436,
            "profileUrl": "https://www.instagram.com/hufureview/",
            "searchTerm": "review mỹ phẩm",
        }
        verified = profile(
            "hufureview",
            "Review Mỹ phẩm và Tips Làm đẹp",
            "Tips làm đẹp và review mỹ phẩm · hufureview@gmail.com",
            category="Reel creator",
            followers=17436,
            captions=["Sản phẩm mới từ @fwee_makeup_vn giúp làn da đẹp và mềm mại với mình"] * 12,
        )
        result = collect_influencers(
            config(target_count=1),
            None,
            apify_token="secret",
            apify_search_fetcher=lambda terms, **kwargs: [discovery],
            apify_hashtag_fetcher=lambda tags, **kwargs: [],
            apify_fetcher=lambda handles, **kwargs: [verified],
        )
        self.assertEqual(len(result["records"]), 1)
        self.assertEqual(result["records"][0]["grade"], "B")
        self.assertEqual(result["records"][0]["apify_country_confidence"], "중간")
        self.assertEqual(result["query_result_counts"]["[Apify Instagram 사용자검색] review mỹ phẩm"], 1)
        self.assertEqual(
            [row["provider"] for row in result["api_response_log"]],
            ["Apify Instagram Search", "Apify Instagram Hashtag"],
        )

    def test_excel_reports_apify_provider_and_actor_totals(self):
        payload = {
            "records": [],
            "review_records": [],
            "rejected_records": [],
            "queries": [
                "[Apify Instagram 사용자검색] review mỹ phẩm",
                "[Apify Instagram 해시태그] #kocvietnam",
            ],
            "query_result_counts": {
                "[Apify Instagram 사용자검색] review mỹ phẩm": 7,
                "[Apify Instagram 해시태그] #kocvietnam": 4,
            },
            "api_response_log": [
                {"provider": "Apify Instagram Search", "query": "사용자 검색어 1개 일괄 실행", "result_count": 7, "http_status": 200},
                {"provider": "Apify Instagram Hashtag", "query": "해시태그 1개 일괄 실행", "result_count": 4, "http_status": 200},
            ],
            "rejected_reasons": {},
            "summary": {},
            "safety": {},
            "config": {"search_provider": "Brave"},
        }
        workbook = load_workbook(io.BytesIO(build_influencer_excel(payload)), data_only=True, read_only=True)
        rows = list(workbook["03_검색근거"].iter_rows(values_only=True))
        workbook.close()
        search_row = next(row for row in rows if row[3] == "[Apify Instagram 사용자검색] review mỹ phẩm")
        hashtag_row = next(row for row in rows if row[3] == "[Apify Instagram 해시태그] #kocvietnam")
        self.assertEqual(search_row[2:5], ("Apify Instagram Search", "[Apify Instagram 사용자검색] review mỹ phẩm", 7))
        self.assertEqual(hashtag_row[2:5], ("Apify Instagram Hashtag", "[Apify Instagram 해시태그] #kocvietnam", 4))


if __name__ == "__main__":
    unittest.main(verbosity=2)
