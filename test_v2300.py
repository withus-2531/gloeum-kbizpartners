import io
import json
import unittest

from openpyxl import load_workbook

from apify_instagram_client import fetch_instagram_hashtag_posts, fetch_instagram_search_profiles
from influencer_core import (
    InfluencerConfig,
    build_instagram_discovery_hashtags,
    build_instagram_discovery_terms,
    collect_influencers,
    discovery_item_to_search_result,
    evaluate_apify_profile,
)
from influencer_exporter import build_influencer_excel
from version import APP_VERSION


def config(**changes):
    values={
        "country":"Vietnam","category":"뷰티·화장품",
        "keywords":["beauty","skincare","makeup"],"target_count":1,
        "max_queries":4,"request_delay":0,"use_apify_verification":True,
        "use_apify_discovery":True,"use_web_fallback":False,
    }
    values.update(changes)
    return InfluencerConfig(**values)


def creator(handle="linhbeauty"):
    return {
        "inputUrl":f"https://www.instagram.com/{handle}/",
        "url":f"https://www.instagram.com/{handle}/",
        "username":handle,"fullName":"Linh Beauty Creator",
        "biography":"Digital creator based in Ho Chi Minh City · beauty and skincare reviews",
        "businessCategoryName":"Digital creator","followersCount":25000,
        "isBusinessAccount":True,
        "latestPosts":[{"caption":"Vietnam skincare serum review and makeup tutorial"} for _ in range(12)],
    }


class FakeResponse:
    def __init__(self,payload): self.payload=json.dumps(payload).encode("utf-8")
    def __enter__(self): return self
    def __exit__(self,*args): return False
    def read(self): return self.payload


class V2300Tests(unittest.TestCase):
    def test_version(self):
        self.assertEqual(APP_VERSION,"2.30.3")

    def test_vietnam_discovery_terms_and_hashtags_are_localized(self):
        terms=build_instagram_discovery_terms(config())
        hashtags=build_instagram_discovery_hashtags(config())
        self.assertIn("review mỹ phẩm",terms)
        self.assertIn("kocvietnam",hashtags)
        self.assertLessEqual(len(terms),10)
        self.assertLessEqual(len(hashtags),8)

    def test_search_actor_uses_official_user_search_input(self):
        captured={}
        def opener(request,timeout):
            captured.update({"url":request.full_url,"body":json.loads(request.data),"auth":request.headers.get("Authorization")})
            return FakeResponse([{"username":"linhbeauty"}])
        rows=fetch_instagram_search_profiles(["Vietnam beauty blogger","review mỹ phẩm"],token="secret",search_limit=12,opener=opener)
        self.assertEqual(len(rows),1)
        self.assertIn("apify~instagram-search-scraper",captured["url"])
        self.assertNotIn("secret",captured["url"])
        self.assertEqual(captured["auth"],"Bearer secret")
        self.assertEqual(captured["body"]["searchType"],"user")
        self.assertEqual(captured["body"]["searchLimit"],12)
        self.assertFalse(captured["body"]["enhanceUserSearchWithFacebookPage"])

    def test_hashtag_actor_uses_bounded_post_input(self):
        captured={}
        def opener(request,timeout):
            captured.update({"url":request.full_url,"body":json.loads(request.data)})
            return FakeResponse([{"ownerUsername":"linhbeauty"}])
        rows=fetch_instagram_hashtag_posts(["#kocvietnam","reviewmypham"],token="secret",results_limit=10,opener=opener)
        self.assertEqual(len(rows),1)
        self.assertIn("apify~instagram-hashtag-scraper",captured["url"])
        self.assertEqual(captured["body"],{"hashtags":["kocvietnam","reviewmypham"],"resultsType":"posts","resultsLimit":10})

    def test_search_and_post_rows_normalize_to_profiles(self):
        user=discovery_item_to_search_result({
            "username":"linhbeauty","fullName":"Linh Nguyen",
            "biography":"Vietnam beauty blogger based in Saigon","followersCount":25000,
            "profileUrl":"https://www.instagram.com/linhbeauty/","searchTerm":"Vietnam beauty blogger",
        },"Apify Instagram 사용자검색")
        post=discovery_item_to_search_result({
            "ownerUsername":"linhbeauty","ownerFullName":"Linh Nguyen",
            "caption":"Review mỹ phẩm tại Ho Chi Minh City","url":"https://www.instagram.com/p/ABC/",
        },"Apify Instagram 해시태그게시물","reviewmypham")
        self.assertEqual(user["url"],"https://www.instagram.com/linhbeauty/")
        self.assertIn("25,000 followers",user["content"])
        self.assertEqual(post["url"],"https://www.instagram.com/linhbeauty/")
        self.assertEqual(post["instagram_discovery_original_url"],"https://www.instagram.com/p/ABC/")

    def test_native_discovery_runs_without_web_provider(self):
        def search_fetcher(terms,**kwargs):
            return [{
                "username":"linhbeauty","fullName":"Linh Beauty Creator",
                "biography":"Digital creator based in Ho Chi Minh City · beauty and skincare reviews",
                "businessCategoryName":"Digital creator","followersCount":25000,
                "profileUrl":"https://www.instagram.com/linhbeauty/","searchTerm":terms[0],
            }]
        result=collect_influencers(
            config(),None,apify_token="secret",
            apify_search_fetcher=search_fetcher,
            apify_hashtag_fetcher=lambda tags,**kwargs:[],
            apify_fetcher=lambda handles,**kwargs:[creator(handle) for handle in handles],
        )
        self.assertEqual([row["handle"] for row in result["records"]],["linhbeauty"])
        self.assertEqual(result["summary"]["instagram_discovery_calls"],2)
        self.assertEqual(result["summary"]["instagram_discovered_handles"],1)
        self.assertEqual(result["summary"]["web_fallback_calls"],0)
        self.assertTrue(result["safety"]["apify_instagram_discovery_used"])

    def test_observed_business_accounts_are_excluded(self):
        for handle in ("beautygarden.vn","swe.vn","sunnyson.vn","kirobeauty","kosallure","kkwbeauty","skincareedit"):
            with self.subTest(handle=handle):
                row=creator(handle)
                result=evaluate_apify_profile(row,config())
                self.assertEqual(result["status"],"제외_기업계정")

    def test_excel_exposes_native_discovery_metrics(self):
        payload={
            "records":[],"review_records":[],"rejected_records":[],"queries":[],
            "api_response_log":[],"rejected_reasons":{},"errors":[],"config":{},
            "summary":{"instagram_discovery_calls":2,"instagram_search_results":80,
                       "instagram_hashtag_posts":40,"instagram_discovered_handles":35,
                       "web_fallback_calls":0},
            "safety":{"apify_instagram_discovery_used":True},
        }
        workbook=load_workbook(io.BytesIO(build_influencer_excel(payload)),data_only=True,read_only=True)
        values=[cell for row in workbook["03_검색근거"].iter_rows(values_only=True) for cell in row]
        workbook.close()
        self.assertIn("Instagram 발견 고유계정",values)
        self.assertIn("웹 보강 호출",values)
        self.assertIn("Apify Instagram 후보 발견",values)


if __name__=="__main__":
    unittest.main(verbosity=2)
